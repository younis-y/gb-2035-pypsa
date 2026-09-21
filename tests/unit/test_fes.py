import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from gb2035.data.fes import FES_METRICS, extract_fes_2035, fes_value, load_fes


def make_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "F.55"
    ws.cell(row=1, column=1, value="F.55: Offshore wind capacity")
    years = [dt.datetime(y, 1, 1) for y in (2024, 2025, 2030, 2035, 2040, 2050)]
    ws.cell(row=8, column=13, value="GW")
    for j, y in enumerate(years):
        ws.cell(row=8, column=14 + j, value=y)
    for i, (label, vals) in enumerate(
        [
            ("History", [15, 15, None, None, None, None]),
            ("Electric Engagement", [15, 16, 47.8, 67.5, 100, 110]),
            ("Holistic Transition", [15, 16, 46.5, 86.1, 120, 130]),
        ]
    ):
        ws.cell(row=9 + i, column=13, value=label)
        for j, v in enumerate(vals):
            ws.cell(row=9 + i, column=14 + j, value=v)
    # second block
    ws.cell(row=18, column=13, value="GW")
    for j, y in enumerate(years):
        ws.cell(row=18, column=14 + j, value=y)
    ws.cell(row=19, column=13, value="Holistic Transition")
    ws.cell(row=19, column=17, value=0.123)
    wb.save(path)


def test_extract_blocks(tmp_path: Path, monkeypatch):
    make_workbook(tmp_path / "fes.xlsx")
    monkeypatch.setattr(
        "gb2035.data.fes.FES_METRICS",
        {"offshore_wind_gw": ("F.55", 0, "GW"), "offshore_second": ("F.55", 1, "GW")},
    )
    out = extract_fes_2035(tmp_path / "fes.xlsx")
    assert fes_value(out, "offshore_wind_gw", "Holistic Transition") == 86.1
    assert fes_value(out, "offshore_wind_gw", "Electric Engagement") == 67.5
    assert fes_value(out, "offshore_second", "Holistic Transition") == 0.123
    assert "History" not in set(out["pathway"])
    assert set(out.columns) == {"metric", "pathway", "value", "unit", "sheet", "block"}


def test_metric_table_has_required_entries():
    for m in [
        "consumer_demand_twh",
        "peak_demand_gw",
        "offshore_wind_gw",
        "onshore_wind_gw",
        "solar_gw",
        "battery_gw",
        "interconnector_gw",
        "nuclear_gw",
        "gas_ccus_gw",
        "hydrogen_generation_gw",
        "unabated_gas_gw",
        "industrial_h2_twh",
    ]:
        assert m in FES_METRICS


@pytest.mark.slow
def test_real_workbook_values(repo_root: Path):
    wb = repo_root / "data" / "raw" / "fes2025_data_workbook.xlsx"
    if not wb.exists():
        pytest.skip("run gb2035 retrieve first")
    out = extract_fes_2035(wb)
    ht = {m: fes_value(out, m, "Holistic Transition") for m in FES_METRICS}
    assert ht["consumer_demand_twh"] == pytest.approx(388.303, abs=0.01)
    assert ht["offshore_wind_gw"] == pytest.approx(86.144, abs=0.01)
    assert ht["nuclear_gw"] == pytest.approx(5.04, abs=0.01)
    assert ht["gas_ccus_gw"] == pytest.approx(8.093, abs=0.01)
    assert ht["hydrogen_generation_gw"] == pytest.approx(2.567, abs=0.01)
    assert ht["industrial_h2_twh"] == pytest.approx(19.79, abs=0.01)
    assert ht["interconnector_gw"] == pytest.approx(19.4, abs=0.01)


def test_committed_csv_matches_workbook_shape(repo_root: Path):
    fes = load_fes(repo_root / "data" / "derived" / "fes2025_gb_2035.csv")
    assert fes_value(fes, "consumer_demand_twh", "Electric Engagement") == pytest.approx(
        407.789, abs=0.01
    )
    assert set(FES_METRICS) <= set(fes["metric"])
